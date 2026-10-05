#!/usr/bin/env python3
"""Generate Crawler Protocol achievement catalog (1000+ rows) as Excel + CSV."""
from __future__ import annotations

import csv
import itertools
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_XLSX = Path("/workspace/crawler-protocol/reference/achievements.xlsx")
OUT_CSV = Path("/workspace/crawler-protocol/reference/achievements.csv")

FLOORS = {
    1: "Orientation Pit",
    2: "Soft Open",
    3: "Kingdom Continuum — Rural",
    4: "Knotline",
    5: "Bubble Farms",
    6: "Kingdom Continuum — Docks",
    7: "Greenlight Gauntlet",
    8: "Remnant Parade",
    9: "Kingdom Continuum — Castle",
    10: "Last Heat Standing",
    11: "Ratings Jubilee",
    12: "Kingdom Continuum — Olympus",
    13: "Ledger Depths",
    14: "Continuity Auction",
    15: "Ash Choir",
    16: "Fracture Gallery",
    17: "Rigging Catwalk",
    18: "Cage Bottom",
}

HEADERS = [
    "ID",
    "Name",
    "Criteria",
    "Floor",
    "Category",
    "Rarity",
    "Has Reward",
    "Reward",
    "Canon Flavor",
    "Notes",
]

rows: list[dict] = []
n = 0


def reward_for(rarity: str, has_reward: bool) -> str:
    if not has_reward:
        return "— (brag / title only)"
    if rarity == "Legendary":
        return "Legendary loot roll + Title + cosmetic flair"
    if rarity == "Rare":
        return "Rare loot roll + Title"
    return "Title + 25 System Credits"


def add(
    name: str,
    criteria: str,
    floor,
    category: str,
    rarity: str = "Common",
    has_reward: bool = True,
    reward: str = "",
    canon: str = "No",
    notes: str = "",
):
    global n
    n += 1
    if has_reward and not reward:
        reward = reward_for(rarity, True)
    if not has_reward:
        reward = "— (brag / title only)"
    rows.append(
        {
            "ID": f"ACH-{n:04d}",
            "Name": name,
            "Criteria": criteria,
            "Floor": str(floor),
            "Category": category,
            "Rarity": rarity,
            "Has Reward": "Yes" if has_reward else "No",
            "Reward": reward,
            "Canon Flavor": canon,
            "Notes": notes,
        }
    )


# ——— Hand-authored globals ———
for name, crit, cat, rarity, hr, reward, canon, notes in [
    ("First Blood", "Deal the killing blow on your first mob.", "Combat", "Common", True, "Title: Blooded + junk weapon oil", "No", ""),
    ("Overkill Artist", "Deal 3x a target's remaining HP in one hit.", "Combat", "Rare", True, "", "No", ""),
    ("Crowd Control King", "Apply a control effect to 5+ enemies in one fight.", "Combat", "Rare", True, "", "No", ""),
    ("David vs Goliath", "Kill a creature 5+ levels above you.", "Combat", "Legendary", True, "", "No", ""),
    ("No Hit Wonder", "Win a fight without taking damage.", "Combat", "Rare", True, "", "No", ""),
    ("Cartographer's Scratch", "Reveal 50% of a floor's fog of war.", "Exploration", "Common", True, "", "No", ""),
    ("Full Fog Breaker", "Reveal 100% of a floor map.", "Exploration", "Legendary", True, "Map cosmetic + rare roll", "No", ""),
    ("Secret Door Sniffer", "Find 3 hidden doors on one floor.", "Exploration", "Rare", True, "", "No", ""),
    ("Silver Tongue", "Succeed on 5 Presence checks in one session.", "Social", "Common", True, "", "No", ""),
    ("Barter Baron", "Haggle a vendor price down by 25%+.", "Social", "Rare", True, "", "No", ""),
    ("Name Dropper", "Invoke your Broadcast Name to gain a social benefit.", "Social", "Rare", True, "", "No", ""),
    ("Pocket Lint", "Loot your first container.", "Loot", "Common", True, "25 credits", "No", ""),
    ("Named Drop", "Obtain a Named item.", "Loot", "Rare", True, "", "No", ""),
    ("Hoarder", "Carry 20+ unique junk items at once.", "Loot", "Common", False, "", "No", ""),
    ("Down But Not Out", "Be stabilized from 0 HP.", "Survival", "Common", True, "", "No", ""),
    ("Close Call", "Survive at 1 HP.", "Survival", "Common", True, "", "No", ""),
    ("Wipe Adjacent", "Be the sole survivor of a near-wipe.", "Survival", "Legendary", True, "", "No", ""),
    ("Buddy System", "Revive or stabilize an ally.", "Party", "Common", True, "", "No", ""),
    ("Combo Meal", "Score a combo kill with an ally on the same turn.", "Party", "Rare", True, "", "No", ""),
    ("Sacrifice Play", "Take a hit that would have downed an ally.", "Party", "Rare", True, "", "No", ""),
    ("Bad Idea Bear", "Propose a plan everyone agrees is terrible — then succeed.", "Humor", "Rare", True, "", "No", ""),
    ("Tutorial Complete", "Finish Floor 1 alive.", "Meta", "Common", True, "Title: Oriented", "No", ""),
    ("Season Ticket", "Attend 5 play sessions.", "Meta", "Common", True, "", "No", ""),
    ("Rules Lawyer Adjacent", "Cite a Core Rules line that changes a ruling.", "Meta", "Rare", True, "Title: Fine Print", "No", ""),
    # Canon-flavored
    ("Saferoom Skeptic", "Enter a saferoom and immediately search it for traps.", "Canon Flavor", "Rare", True, "Trap kit upgrade", "Yes", "saferoom paranoia trope"),
    ("Cat Person", "Befriend or recruit a feline (or feline-coded) companion.", "Canon Flavor", "Legendary", True, "Companion slot perk", "Yes", "iconic companion energy"),
    ("Sponsorship Bite", "Accept a sponsor deal with a nasty rider.", "Canon Flavor", "Rare", True, "Sponsor currency + debuff", "Yes", "sponsor economy"),
    ("Sponsorship Spit", "Refuse a sponsor deal on principle.", "Canon Flavor", "Rare", False, "", "Yes", "sponsor economy"),
    ("Desperado Inventory", "Win a fight using only improvised/junk weapons.", "Canon Flavor", "Rare", True, "Named junk weapon", "Yes", "improvised combat"),
    ("The Audience Is Watching", "Trigger a ratings spike event (AI DM flags it).", "Canon Flavor", "Rare", True, "Ratings buff (1 floor)", "Yes", "show framing"),
    ("Dungeon Anarchist Adjacent", "Break a floor rule in a way that helps the party and isn't hard-punished.", "Canon Flavor", "Legendary", True, "Legendary utility", "Yes", "rule-bending spirit"),
    ("Achievement for Achievements", "Earn 25 achievements in one campaign.", "Canon Flavor", "Common", True, "Title: Collector", "Yes", "meta game-show"),
    ("Patch Notes Incoming", "Survive a mid-floor rule change.", "Canon Flavor", "Rare", True, "", "Yes", "show-trope"),
    ("Loot Council", "Let the party vote on a Named drop.", "Canon Flavor", "Common", True, "", "Yes", ""),
    ("Murder-Show Smile", "Deliver a one-liner after a kill that the table quotes later.", "Canon Flavor", "Common", False, "", "Yes", "dark comedy"),
]:
    add(name, crit, "Any", cat, rarity, hr, reward, canon, notes)

# Continuum ladder (3/6/9/12)
CONTINUUM = [
    (3, "First Village Firelight", "Spend a night (rest) in a rural settlement saferoom.", "Continuum", "Common", True, "Title: Guest"),
    (3, "Millroad Regular", "Visit the same village landmark twice.", "Continuum", "Common", True, ""),
    (3, "Harvest Helper", "Complete a farmstead or harvest quest.", "Continuum", "Common", True, "Food supplies"),
    (3, "Low Chapel Candle", "Make an offering at a rural shrine.", "Continuum", "Common", True, "Minor blessing (1 fight)"),
    (3, "Settlement Defender", "Save a village from a raid or mob wave.", "Continuum", "Rare", True, "Rare roll + village favor"),
    (3, "Gossip Network", "Learn 5 distinct rumors in rural zones.", "Continuum", "Common", True, ""),
    (3, "Dirt Road Diplomat", "End a feud between two settlements peacefully.", "Continuum", "Rare", True, ""),
    (3, "Broadcast Name: Local Legend", "Hear NPCs use your Broadcast Name unprompted.", "Continuum", "Rare", True, "Social buff in Continuum"),
    (3, "Odd Coin Curiosity", "Show the Odd Coin to a village elder or vendor.", "Story Seed", "Common", True, "Lore unlock"),
    (3, "Benefactor Folk Omen", "Have an NPC interpret the Benefactor Signal as superstition.", "Story Seed", "Common", False, ""),
    (6, "Harbor Arrival", "Enter the docks district for the first time.", "Continuum", "Common", True, ""),
    (6, "Customs Circumvented", "Smuggle an item past dock customs.", "Continuum", "Rare", True, ""),
    (6, "Sponsor Skybox Wave", "Be acknowledged from a sponsor skybox.", "Continuum", "Rare", True, "Sponsor currency"),
    (6, "Hunter License Seen", "Witness or steal a hunter license/badge.", "Continuum", "Rare", True, "Hunter Badge progress"),
    (6, "Pier Fight", "Win a combat on the docks without falling in the water.", "Combat", "Common", True, ""),
    (6, "Man Overboard", "Fall (or push an enemy) into dock water.", "Humor", "Common", False, ""),
    (6, "Refugee Reunion", "Meet a Floor 3 NPC again at the outskirts/docks.", "Continuum", "Legendary", True, "Legendary social reward"),
    (6, "Warehouse Sweep", "Clear a warehouse of hostiles and loot it fully.", "Exploration", "Common", True, ""),
    (6, "Benefactor Hard Choice", "Accept or refuse the Benefactor's midgame offer.", "Story Seed", "Legendary", True, "Allegiance flag"),
    (6, "Sponsor Bait Taken", "Sign a sponsor deal at the docks.", "Canon Flavor", "Rare", True, "Deal rider logged"),
    (9, "Castle Gates", "Enter the castle grounds.", "Continuum", "Common", True, "Title: Court-Touched"),
    (9, "Upper City Stroll", "Visit 3 noble-district landmarks.", "Exploration", "Common", True, ""),
    (9, "Mandate Accepted", "Accept an officer/faction mandate.", "Continuum", "Rare", True, "Faction Flag progress"),
    (9, "Mandate Burned", "Publicly refuse a mandate.", "Continuum", "Rare", False, ""),
    (9, "Throne Room Adjacent", "Reach the castle throne room (or equivalent).", "Continuum", "Rare", True, ""),
    (9, "Court Etiquette", "Pass a formal etiquette / Presence challenge at court.", "Social", "Common", True, ""),
    (9, "Dock Money at Court", "Use a Floor 6 sponsor connection inside the castle.", "Continuum", "Legendary", True, ""),
    (9, "Village Saint Elevated", "Discover a Floor 3 shrine figure now venerated at court.", "Continuum", "Rare", True, "Lore + minor buff"),
    (9, "Faction Flag Claimed", "Plant or claim a Faction Flag objective.", "Story Seed", "Legendary", True, "Flag item"),
    (9, "Assassination Foiled", "Stop an attempt on a key court NPC (or yourself).", "Combat", "Rare", True, ""),
    (12, "Olympus Ascent", "Reach the deity tier of the Kingdom Continuum.", "Continuum", "Common", True, "Title: Skybound"),
    (12, "Pantheon Audience", "Speak before a deity or pantheon court.", "Continuum", "Rare", True, ""),
    (12, "Throne Vow Spoken", "Swear (or refuse) the Throne Vow.", "Story Seed", "Legendary", True, "Vow flag"),
    (12, "Village God Made Real", "Confront the divine form of a Floor 3 folk saint/god.", "Continuum", "Legendary", True, "Divine relic roll"),
    (12, "Castle Echo", "See Floor 9 castle architecture mirrored in Olympus.", "Exploration", "Common", True, ""),
    (12, "Dock Idol Awakened", "A Floor 6 harbor idol appears as a real divine actor.", "Continuum", "Rare", True, ""),
    (12, "Benefactor at the Peak", "The Benefactor thread intersects pantheon politics.", "Story Seed", "Legendary", True, "Endgame allegiance options"),
    (12, "Divine Etiquette", "Survive a deity's test of manners or wit.", "Social", "Rare", True, ""),
    (12, "Vowbreaker Thunder", "Break a divine vow and live to the stairwell.", "Continuum", "Legendary", True, "Cursed legendary loot"),
    (12, "Last Hometown Light", "Say goodbye to a Continuum NPC who will not return.", "Social", "Rare", False, ""),
]
for fl, name, crit, cat, rarity, hr, reward in CONTINUUM:
    canon = "Yes" if cat in ("Story Seed", "Canon Flavor") or "Benefactor" in name or "Odd Coin" in name or "Sponsor" in name else "No"
    add(name, crit, fl, cat, rarity, hr, reward, canon, "continuum ladder")

# Floor-specific flavor
FLAVOR = {
    1: [
        ("Orientation Graduate", "Exit Floor 1 through the stairwell alive.", "Meta", "Common", True, "Title: Oriented", "No"),
        ("Vending Gore", "Loot a vending machine on Floor 1.", "Loot", "Common", True, "", "No"),
        ("Tutorial Elite", "Defeat the Orientation elite.", "Combat", "Common", True, "", "No"),
        ("Don't Skip Intro", "Read/listen to a full System tutorial popup.", "Meta", "Common", False, "", "Yes"),
        ("Crowbar Starter", "Equip a junk melee as your first weapon.", "Loot", "Common", True, "", "Yes"),
        ("Odd Coin Pocketed", "Pick up the Odd Coin.", "Story Seed", "Common", True, "Seed flagged", "Yes"),
        ("Static in the Ear", "Trigger the Benefactor Signal static event.", "Story Seed", "Common", True, "Seed flagged", "Yes"),
        ("Corpse Lesson", "Loot a fallen crawler NPC tutorial corpse.", "Loot", "Common", False, "", "Yes"),
        ("Stair Bait Resisted", "Ignore the first obvious stair bait trap.", "Exploration", "Rare", True, "", "No"),
        ("First Down", "Drop to 0 HP on Floor 1.", "Survival", "Common", False, "", "No"),
    ],
    2: [
        ("Soft Open Survivor", "Clear Floor 2.", "Meta", "Common", True, "", "No"),
        ("Trap Sense Tingling", "Disarm 3 traps on Floor 2.", "Exploration", "Rare", True, "", "No"),
        ("Saferoom With Teeth", "Discover a saferoom hazard on Floor 2.", "Survival", "Common", True, "", "Yes"),
        ("Companion Spark", "Unlock or recruit a companion slot option.", "Party", "Rare", True, "", "No"),
        ("Guide Contract Glimpse", "Encounter the Guide Contract seed.", "Story Seed", "Common", True, "", "Yes"),
        ("Maintenance Maze", "Clear the maintenance/side wing.", "Exploration", "Common", True, "", "No"),
    ],
    4: [
        ("Transfer Complete", "Ride or solve a Knotline transfer.", "Exploration", "Common", True, "", "No"),
        ("Rail Key Found", "Obtain a transit key.", "Story Seed", "Rare", True, "", "Yes"),
        ("Station Master", "Secure a Knotline station hub.", "Combat", "Rare", True, "", "No"),
        ("Wrong Platform", "Board the incorrect line and survive.", "Humor", "Common", False, "", "No"),
        ("Relay Runner", "Complete a team relay objective.", "Party", "Rare", True, "", "No"),
        ("Geometry Nosebleed", "Fail a Wits check due to non-Euclidean space and recover.", "Survival", "Common", False, "", "No"),
    ],
    5: [
        ("Bubble Hopper", "Enter 3 different biome bubbles.", "Exploration", "Common", True, "", "No"),
        ("Core Sample", "Obtain a Bubble Core sample.", "Story Seed", "Rare", True, "", "Yes"),
        ("Coop Mandatory", "Complete an objective requiring temporary allies.", "Party", "Rare", True, "", "No"),
        ("Diplomatic Bubble", "Broker a deal between two bubble factions.", "Social", "Rare", True, "", "No"),
        ("Weather Whiplash", "Suffer two opposite biome hazards in one session.", "Survival", "Common", False, "", "No"),
        ("Biome Tourist", "Take a souvenir from each visited bubble.", "Loot", "Common", True, "", "No"),
    ],
    7: [
        ("Greenlight Go", "Start a timed heat/gauntlet run.", "Speed", "Common", True, "", "No"),
        ("Par Time", "Beat the floor par time.", "Speed", "Rare", True, "", "No"),
        ("Shortcut Graffiti", "Find and use a marked shortcut.", "Exploration", "Rare", True, "", "Yes"),
        ("Zero-Death Heat", "Complete a timed segment with no downs.", "Survival", "Legendary", True, "", "No"),
        ("Map Steal", "Steal or copy a route map from a rival.", "Story Seed", "Rare", True, "", "Yes"),
        ("Wrong Way Faster", "Win via a route everyone thought was wrong.", "Humor", "Rare", True, "", "No"),
    ],
    8: [
        ("Echo Catcher", "Capture or bind an echo card/unit.", "Story Seed", "Rare", True, "", "Yes"),
        ("Self Duel", "Fight an echo of a party member (or yourself).", "Combat", "Legendary", True, "", "Yes"),
        ("Parade Route", "Survive the Remnant Parade path.", "Survival", "Common", True, "", "No"),
        ("Deck Synergy", "Win a fight using 3 echo abilities together.", "Combat", "Rare", True, "", "No"),
        ("Possession Resist", "Resist a possession attempt.", "Survival", "Rare", True, "", "No"),
        ("Folklore Zone", "Enter a folklore-marked echo zone.", "Exploration", "Common", True, "", "No"),
    ],
    10: [
        ("Garage Band", "Assign crew roles in a heat garage.", "Party", "Common", True, "", "No"),
        ("Heat Win", "Win a ranked heat.", "Speed", "Rare", True, "", "No"),
        ("Not Last", "Avoid last place in an elimination heat.", "Survival", "Common", True, "", "Yes"),
        ("Black Box Ping", "Recover or detect a Black Box.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Pit Stop", "Repair during a heat without being eliminated.", "Crafting", "Common", True, "", "No"),
        ("Wrecking Ball", "Cause a rival vehicle/runner wipe.", "Combat", "Rare", True, "", "No"),
    ],
    11: [
        ("Jubilee Guest", "Enter the Ratings Jubilee grounds.", "Meta", "Common", True, "", "No"),
        ("Crowd Turn", "Flip audience favor during an arena bout.", "Social", "Rare", True, "", "Yes"),
        ("Named Duel", "Win a named opponent duel.", "Combat", "Rare", True, "", "No"),
        ("VIP Ticket", "Obtain a Jubilee Ticket / VIP pass.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Parade Float", "Ride or sabotage a parade element.", "Humor", "Common", True, "", "No"),
        ("Spectacle Finish", "End a fight with a showy finisher the crowd loves.", "Combat", "Common", False, "", "Yes"),
    ],
    13: [
        ("Audit Opened", "Begin a ledger audit path.", "Exploration", "Common", True, "", "No"),
        ("Name Restored", "Restore an erased name (ally/NPC/self).", "Story Seed", "Legendary", True, "", "Yes"),
        ("Quiet Horror", "Traverse a silent vault without triggering alarms.", "Exploration", "Rare", True, "", "No"),
        ("Debt Notice", "Receive a System debt notice and resolve or defer it.", "Meta", "Common", True, "", "No"),
        ("Filing Cathedral", "Reach the deepest filing hall.", "Exploration", "Rare", True, "", "No"),
        ("Echo-Self Paperwork", "Find ledger records about your own echo/self.", "Story Seed", "Rare", True, "", "Yes"),
    ],
    14: [
        ("Auction Paddle", "Place a bid in the Continuity Auction.", "Social", "Common", True, "", "No"),
        ("Lot Won", "Win a continuity lot.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Null Currency", "Earn or spend null/continuity currency.", "Loot", "Rare", True, "", "No"),
        ("Evidence Entered", "Submit Black Box evidence into an auction.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Outbid Trauma", "Lose a lot you desperately wanted.", "Humor", "Common", False, "", "No"),
        ("Rewrite Token", "Obtain a rewrite token.", "Loot", "Legendary", True, "", "No"),
    ],
    15: [
        ("Ash on the Tongue", "Enter the Ash Choir layer.", "Meta", "Common", True, "", "No"),
        ("Pact Offered", "Receive an infernal pact offer.", "Story Seed", "Common", True, "", "Yes"),
        ("Pact Refused", "Refuse an infernal pact.", "Story Seed", "Rare", False, "", "Yes"),
        ("Pact Signed", "Sign an infernal pact (rider logged).", "Story Seed", "Legendary", True, "Power + cost", "Yes"),
        ("Choir Silence", "Defeat or silence a choir boss.", "Combat", "Legendary", True, "", "No"),
        ("Vow Invoice", "Have a prior Throne Vow come due here.", "Story Seed", "Legendary", True, "", "Yes"),
    ],
    16: [
        ("Mirror Entry", "Enter the Fracture Gallery.", "Meta", "Common", True, "", "No"),
        ("Mirror Kill", "Defeat a mirror-self or clone hostile.", "Combat", "Rare", True, "", "Yes"),
        ("Self Reunite", "Reintegrate a split identity safely.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Clone Law Briefing", "Learn and cite a clone-law rule in play.", "Meta", "Common", True, "", "No"),
        ("Gallery Clear", "Clear the primary gallery wing.", "Exploration", "Rare", True, "", "No"),
        ("Sold Face", "Discover a market selling your likeness.", "Social", "Rare", False, "", "Yes"),
    ],
    17: [
        ("Catwalk Legs", "Traverse the Rigging Catwalk without falling.", "Survival", "Common", True, "", "No"),
        ("Stagehand Tools", "Obtain a stagehand tool item.", "Loot", "Common", True, "", "No"),
        ("Broadcast Hack", "Successfully hack a broadcast node.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Black Box Air", "Transmit Black Box data on the rigging.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Lighting Grid Tour", "Visit rigging overlooks of 3 prior biomes.", "Exploration", "Rare", True, "", "No"),
        ("Escape Route Mapped", "Map an escape route for Floor 18.", "Exploration", "Rare", True, "", "No"),
    ],
    18: [
        ("Cage Bottom Tickets", "Enter the Cage Bottom.", "Meta", "Common", True, "Title: Finalist", "No"),
        ("Owner's Lounge", "Reach the owner's lounge / VIP cage area.", "Exploration", "Rare", True, "", "No"),
        ("Ending: Crown Content", "Trigger the Crown Content ending path.", "Meta", "Legendary", True, "Ending flag", "No"),
        ("Ending: Patron Exit", "Trigger the Patron Exit ending path.", "Meta", "Legendary", True, "Ending flag", "Yes"),
        ("Ending: Burn the Feed", "Trigger Burn the Feed using Black Box + Rigging.", "Meta", "Legendary", True, "Ending flag", "Yes"),
        ("Ending: Ledger Ghost", "Trigger the Ledger Ghost ending path.", "Meta", "Legendary", True, "Ending flag", "No"),
        ("Ending: New Showrunners", "Trigger the New Showrunners ending path.", "Meta", "Legendary", True, "Ending flag", "Yes"),
        ("Spine Convergence", "Have 3+ spine objects matter in the finale scene.", "Story Seed", "Legendary", True, "", "Yes"),
        ("Final One-Liner", "Deliver the campaign's closing line (table votes).", "Humor", "Common", False, "", "Yes"),
        ("Credits Roll", "Complete the campaign session that ends Season Zero.", "Meta", "Legendary", True, "Season cosmetic pack", "No"),
    ],
}
for fl, pack in FLAVOR.items():
    for name, crit, cat, rarity, hr, reward, canon in pack:
        add(name, crit, fl, cat, rarity, hr, reward, canon, "")

# Per-floor template grids
COMBAT = [
    ("Clean Sweep", "Defeat every marked encounter in a designated zone.", "Rare"),
    ("Elite Hunter", "Defeat 3 elite enemies.", "Rare"),
    ("Boss Claim", "Defeat the floor's primary boss or gate guardian.", "Rare"),
    ("Minion Mower", "Defeat 25 standard enemies.", "Common"),
    ("Status Stacker", "Apply 3 different status effects in one fight.", "Common"),
    ("Environmentalist", "Kill an enemy using a floor hazard.", "Rare"),
    ("Perfect Ambush", "Start hidden and win without an ally dropping.", "Rare"),
    ("Last Hit League", "Score 5 killing blows.", "Common"),
]
EXPLORE = [
    ("Zone Tourist", "Enter all primary named zones.", "Common"),
    ("Hidden Cache", "Find a hidden loot cache.", "Common"),
    ("Landmark Check-in", "Interact with the signature landmark.", "Common"),
    ("Side Path", "Discover an optional side path.", "Common"),
    ("Stair Seeker", "Locate the stairwell/gate before defeating the boss.", "Rare"),
    ("Fog Half", "Reveal at least 50% of the map.", "Common"),
    ("Fog Full", "Reveal the entire map.", "Legendary"),
    ("Locksmith", "Open 5 locked containers.", "Common"),
]
SOCIAL = [
    ("Local Contact", "Gain a named NPC ally.", "Common"),
    ("Quest Complete", "Complete a side quest.", "Common"),
    ("Quest Chain", "Complete 3 quests.", "Rare"),
    ("Vendor Regular", "Buy from the same vendor twice.", "Common"),
    ("Rumor Mill", "Collect 3 rumors.", "Common"),
    ("Peaceful Path", "Resolve a hostile encounter via talk.", "Rare"),
]
LOOT = [
    ("Floor Jackpot", "Obtain a Rare or better item.", "Common"),
    ("Named on Site", "Obtain a Named item.", "Rare"),
    ("Junk Architect", "Loot 15 junk items.", "Common"),
    ("Full Sellout", "Sell 10 items to vendors.", "Common"),
    ("Upgrade Desk", "Upgrade or enchant a gear piece.", "Rare"),
]
SURVIVAL = [
    ("No Wipe", "Clear without a party wipe.", "Common"),
    ("Ironman Cub", "Clear without anyone dropping to 0 HP.", "Legendary"),
    ("Bleed Tax", "Spend 3 healing items.", "Common"),
    ("Poison Tolerant", "Suffer and cure a poison/disease.", "Common"),
    ("Timer Terror", "Succeed on a timed objective (if any).", "Rare"),
]
SPEED = [
    ("Speed Clear", "Reach the stairwell under the AI DM's par time.", "Rare"),
    ("Slow Burn", "Take the long route on purpose and still clear.", "Common"),
]
CRAFT = [
    ("Field Craft", "Craft or cook once.", "Common"),
    ("Kit Bash", "Combine two junk items into something useful.", "Rare"),
]
HUMOR = [
    ("Floor Bit", "Create a running joke the table keeps using.", "Common"),
    ("Scenic Panic", "Panic in chat before a fight you then win.", "Common"),
]

for fl, fname in FLOORS.items():
    kind = "continuum" if fl in (3, 6, 9, 12) else ("tutorial" if fl == 1 else "unique")
    for label, crit, rarity in COMBAT:
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Combat", rarity, True, "", "No", kind)
    for label, crit, rarity in EXPLORE:
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Exploration", rarity, True, "", "No", kind)
    for label, crit, rarity in SOCIAL:
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Social", rarity, True, "", "No", kind)
    for label, crit, rarity in LOOT:
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Loot", rarity, True, "", "No", kind)
    for label, crit, rarity in SURVIVAL:
        hr = label != "Bleed Tax"
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Survival", rarity, hr, "", "No", kind)
    for label, crit, rarity in SPEED:
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Speed", rarity, True, "", "No", kind)
    for label, crit, rarity in CRAFT:
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Crafting", rarity, True, "", "No", kind)
    for label, crit, rarity in HUMOR:
        add(f"{fname}: {label}", f"[Floor {fl}] {crit}", fl, "Humor", rarity, False, "", "No", kind)

# Extra action variants to push well over 1000
ACTIONS = [
    ("Open a marked red door", "Exploration"),
    ("Open a marked blue door", "Exploration"),
    ("Survive a night cycle", "Survival"),
    ("Survive a day cycle", "Survival"),
    ("Use a consumable in combat", "Combat"),
    ("Use a consumable out of combat", "Survival"),
    ("Pick a lock successfully", "Exploration"),
    ("Fail a lockpick", "Exploration"),
    ("Win initiative", "Combat"),
    ("Lose initiative and still win the fight", "Combat"),
    ("Deal fire damage", "Combat"),
    ("Deal cold damage", "Combat"),
    ("Deal lightning damage", "Combat"),
    ("Deal poison damage", "Combat"),
    ("Deal psychic damage", "Combat"),
    ("Block a hit with a shield or brace", "Combat"),
    ("Dodge a telegraphed AoE", "Combat"),
    ("Rescue a civilian NPC", "Social"),
    ("Fail to rescue a civilian NPC", "Social"),
    ("Find a journal or lore scrap", "Exploration"),
    ("Read a lore scrap aloud", "Social"),
    ("Destroy a camera or drone", "Combat"),
    ("Hack a terminal successfully", "Exploration"),
    ("Fail a hack", "Exploration"),
    ("Plant a bomb or trap", "Combat"),
    ("Disarm an enemy trap with junk", "Crafting"),
    ("Win a drinking or eating contest", "Humor"),
    ("Lose a betting game", "Humor"),
    ("Win a betting game", "Humor"),
    ("Tame or bribe a beast", "Social"),
    ("Ride a mount or vehicle", "Speed"),
    ("Crash a mount or vehicle", "Humor"),
    ("Cross hazardous water", "Exploration"),
    ("Climb a vertical route", "Exploration"),
    ("Fall and survive", "Survival"),
    ("Use a flare or signal", "Exploration"),
    ("Send a message via System mail", "Meta"),
    ("Receive a threatening System mail", "Meta"),
    ("Equip a full armor set", "Loot"),
    ("Fight bare-handed", "Combat"),
    ("Dual wield in a fight", "Combat"),
    ("Use a two-handed weapon", "Combat"),
    ("Land a critical hit", "Combat"),
    ("Suffer a critical failure", "Humor"),
    ("Stabilize two allies", "Party"),
    ("Be stabilized twice", "Survival"),
    ("Refuse healing out of pride once", "Humor"),
    ("Share food with the party", "Party"),
    ("Steal from a vendor and get away", "Loot"),
    ("Get caught stealing", "Social"),
    ("Pay a fine", "Social"),
    ("Bribe a guard", "Social"),
    ("Tell a lie that becomes true later", "Story Seed"),
    ("Tell the truth when a lie was safer", "Social"),
]

rarities = itertools.cycle(["Common", "Common", "Common", "Rare", "Common", "Rare", "Legendary"])
for fl, fname in FLOORS.items():
    for i, (action, cat) in enumerate(ACTIONS):
        rarity = next(rarities)
        hr = not (i % 7 == 0 and rarity == "Common")
        add(
            f"{fname}: {action}",
            f"[Floor {fl}] {action[0].upper() + action[1:]} (once on this floor).",
            fl,
            cat,
            rarity,
            hr,
            "",
            "No",
            "generated variant",
        )

# Deduplicate names
seen: dict[str, int] = {}
for r in rows:
    name = r["Name"]
    if name in seen:
        seen[name] += 1
        r["Name"] = f"{name} ({seen[name]})"
    else:
        seen[name] = 1

assert len(rows) >= 1000, f"only {len(rows)}"

# Write Excel
wb = Workbook()
summary = wb.active
summary.title = "Summary"
summary["A1"] = "Crawler Protocol — Achievement Catalog"
summary["A1"].font = Font(bold=True, size=14)
summary["A2"] = f"Total achievements: {len(rows)}"
summary["A3"] = f"With rewards: {sum(1 for r in rows if r['Has Reward']=='Yes')}"
summary["A4"] = f"Brag-only (no reward): {sum(1 for r in rows if r['Has Reward']=='No')}"
summary["A5"] = f"Canon flavor tagged: {sum(1 for r in rows if r['Canon Flavor']=='Yes')}"
summary["A7"] = "Rarity"
summary["B7"] = "Count"
summary["A8"] = "Common"
summary["B8"] = sum(1 for r in rows if r["Rarity"] == "Common")
summary["A9"] = "Rare"
summary["B9"] = sum(1 for r in rows if r["Rarity"] == "Rare")
summary["A10"] = "Legendary"
summary["B10"] = sum(1 for r in rows if r["Rarity"] == "Legendary")
summary["A12"] = "Open the Achievements sheet and use Excel filters. Narrative docs stay in markdown."
summary["A13"] = "Draft v1 — rename/expand freely."

ws = wb.create_sheet("Achievements")
fill = PatternFill("solid", fgColor="1F4E79")
font = Font(color="FFFFFF", bold=True)
for col, h in enumerate(HEADERS, 1):
    c = ws.cell(1, col, h)
    c.fill = fill
    c.font = font
    c.alignment = Alignment(horizontal="center")
for i, r in enumerate(rows, 2):
    for col, h in enumerate(HEADERS, 1):
        ws.cell(i, col, r[h])
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows)+1}"
ws.freeze_panes = "A2"
for i, w in enumerate([12, 42, 72, 12, 14, 12, 12, 40, 14, 24], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(OUT_XLSX)
with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=HEADERS)
    w.writeheader()
    w.writerows(rows)

print(len(rows))
print(OUT_XLSX)
print("rewards", sum(1 for r in rows if r["Has Reward"] == "Yes"))
print("canon", sum(1 for r in rows if r["Canon Flavor"] == "Yes"))
print("common", sum(1 for r in rows if r["Rarity"] == "Common"))
print("rare", sum(1 for r in rows if r["Rarity"] == "Rare"))
print("legendary", sum(1 for r in rows if r["Rarity"] == "Legendary"))
